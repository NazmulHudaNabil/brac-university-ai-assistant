import os
import json
from langchain_community.document_loaders import (
    BSHTMLLoader,
    PyPDFLoader,
    TextLoader
)
from langchain_core.documents import Document
import docx
from pptx import Presentation

class DocxLoader:
    def __init__(self, file_path):
        self.file_path = file_path
    
    def load(self):
        doc = docx.Document(self.file_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return [Document(page_content=text, metadata={"source": self.file_path})]

class PptxLoader:
    def __init__(self, file_path):
        self.file_path = file_path
    
    def load(self):
        prs = Presentation(self.file_path)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return [Document(page_content=text, metadata={"source": self.file_path})]

def get_loader(file_path: str, document_type: str):
    """
    Returns the appropriate document loader based on the document type.
    """
    if document_type == "html":
        return BSHTMLLoader(file_path)
    elif document_type == "pdf":
        return PyPDFLoader(file_path)
    elif document_type == "docx":
        return DocxLoader(file_path)
    elif document_type == "pptx":
        return PptxLoader(file_path)
    elif document_type == "txt":
        return TextLoader(file_path, encoding="utf-8")
    else:
        # Fallback to TextLoader
        return TextLoader(file_path, encoding="utf-8")

def load_document(metadata_path: str) -> list[Document]:
    """
    Reads a metadata JSON file to find the companion document, 
    loads its content, and injects metadata into the resulting Documents.
    """
    with open(metadata_path, 'r', encoding='utf-8') as f:
        metadata = json.load(f)
        
    doc_id = metadata.get("document_id")
    doc_type = metadata.get("document_type")
    
    # The actual file has the same name as the JSON file but with the doc_type extension
    base_path = metadata_path.rsplit('.', 1)[0]
    file_path = f"{base_path}.{doc_type}"
    
    if not os.path.exists(file_path):
        print(f"Warning: Data file not found for {metadata_path}")
        return []
        
    loader = get_loader(file_path, doc_type)
    try:
        documents = loader.load()
        
        # Inject our custom metadata into every page/document returned
        for doc in documents:
            doc.metadata["document_id"] = metadata.get("document_id")
            doc.metadata["title"] = metadata.get("title")
            doc.metadata["source_url"] = metadata.get("source_url")
            doc.metadata["document_type"] = metadata.get("document_type")
            
        return documents
    except Exception as e:
        print(f"Error loading {file_path}: {e}")
        return []
